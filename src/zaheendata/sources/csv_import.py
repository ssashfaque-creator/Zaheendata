"""Import spreadsheets: government bulk downloads (licence rolls, Secretary of State extracts),
chamber-of-commerce or association member lists, or your own sheets. Column names are auto-mapped
from a synonym table; pass an explicit mapping to override. Works with CSV and XLSX."""

from __future__ import annotations

import csv
import re
from collections.abc import AsyncIterator, Iterable
from pathlib import Path
from typing import Any, ClassVar

from ..models import Address, GeoPoint, RawRecord, SourceKind, SourceSearchSpec
from .base import Source

SYNONYMS: dict[str, list[str]] = {
    "name": [
        "company",
        "company name",
        "business name",
        "name",
        "organization",
        "account name",
        "dba",
        "legal name",
    ],
    "phone": [
        "phone",
        "phone number",
        "telephone",
        "company phone",
        "primary phone",
        "business phone",
        "hq phone",
    ],
    "website": ["website", "web site", "url", "web address", "domain", "company website", "website url"],
    "line1": ["address", "address line 1", "street", "street address", "address1", "company street address"],
    "city": ["city", "company city", "town"],
    "region": ["state", "st", "province", "company state", "state/province", "region"],
    "postal_code": ["zip", "zip code", "postal code", "postcode", "company zip code"],
    "country": ["country", "company country"],
    "naics": ["naics", "naics code", "primary naics", "primary naics code", "naics 2017", "naics 2022"],
    "employee_count": [
        "employees",
        "employee count",
        "# employees",
        "number of employees",
        "location employees",
        "employees (single site)",
        "estimated num employees",
        "employees here",
    ],
    "employee_range": ["employee range", "employee size", "employees range", "# of employees range"],
    "owner_name": [
        "owner",
        "owner name",
        "contact name",
        "executive name",
        "contact",
        "principal",
        "full name",
    ],
    "owner_first": ["first name", "contact first name", "executive first name"],
    "owner_last": ["last name", "contact last name", "executive last name"],
    "owner_title": ["title", "job title", "contact title", "executive title"],
    "owner_email": ["email", "email address", "contact email", "work email"],
    "owner_phone": ["direct phone", "mobile phone", "contact phone", "direct dial", "mobile"],
    "linkedin_url": ["linkedin", "linkedin url", "person linkedin url", "company linkedin url"],
    "license_number": ["license", "license number", "license no", "licence number", "lic no", "license #"],
    "license_status": ["license status", "status", "licence status"],
    "lat": ["latitude", "lat"],
    "lng": ["longitude", "lng", "lon", "long"],
    "rating": ["rating", "google rating", "stars"],
    "review_count": ["reviews", "review count", "google reviews", "number of reviews", "user ratings total"],
    "source_id": [
        "id",
        "record id",
        "infogroup id",
        "duns",
        "duns number",
        "uei",
        "entity id",
        "loan number",
    ],
    "categories": ["category", "categories", "industry", "primary industry", "business type"],
    "franchise": ["franchise", "franchise flag", "is franchise", "franchise/specialty"],
}


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9#/ ]+", "", header.strip().lower())


def auto_map(headers: Iterable[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    normalized = {_norm(h): h for h in headers}
    for field, names in SYNONYMS.items():
        for n in names:
            if n in normalized and field not in mapping:
                mapping[field] = normalized[n]
                break
    return mapping


def read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(h) if h is not None else "" for h in next(rows)]
        return headers, [dict(zip(headers, r, strict=False)) for r in rows if any(c is not None for c in r)]
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return list(reader.fieldnames or []), list(reader)


def _to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    m = re.search(r"\d[\d,]*", str(value))
    return int(m.group(0).replace(",", "")) if m else None


def _to_float(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def row_to_record(
    row: dict[str, Any],
    mapping: dict[str, str],
    source: str,
    naics: str | None,
    kind: SourceKind = SourceKind.IMPORT,
) -> RawRecord | None:
    def get(field: str) -> Any:
        col = mapping.get(field)
        v = row.get(col) if col else None
        return v.strip() if isinstance(v, str) else v

    name = get("name")
    if not name:
        return None
    owner_name = get("owner_name") or " ".join(filter(None, [get("owner_first"), get("owner_last")])) or None
    lat, lng = _to_float(get("lat")), _to_float(get("lng"))
    categories_raw = get("categories")
    return RawRecord(
        source=source,
        source_kind=kind,
        source_id=str(get("source_id")) if get("source_id") else None,
        name=str(name),
        phone=str(get("phone")) if get("phone") else None,
        website=str(get("website")) if get("website") else None,
        address=Address(
            line1=get("line1"),
            city=get("city"),
            region=get("region"),
            postal_code=str(get("postal_code")) if get("postal_code") else None,
            country=get("country") or "US",
        ),
        location=GeoPoint(lat=lat, lng=lng) if lat is not None and lng is not None else None,
        categories=[c.strip() for c in str(categories_raw).split(";") if c.strip()] if categories_raw else [],
        naics=str(get("naics")) if get("naics") else naics,
        employee_count=_to_int(get("employee_count")),
        employee_range=str(get("employee_range")) if get("employee_range") else None,
        rating=_to_float(get("rating")),
        review_count=_to_int(get("review_count")),
        owner_name=owner_name,
        owner_title=get("owner_title"),
        owner_email=get("owner_email"),
        owner_phone=str(get("owner_phone")) if get("owner_phone") else None,
        linkedin_url=get("linkedin_url"),
        license_number=str(get("license_number")) if get("license_number") else None,
        license_status=get("license_status"),
        attributes={
            "franchise_flag": get("franchise"),
            "import_row": {k: v for k, v in row.items() if v not in (None, "")},
        },
    )


class CsvImportSource(Source):
    name = "csv_import"
    kind = SourceKind.IMPORT
    tags: ClassVar[set[str]] = {"import", "csv", "xlsx"}
    description = (
        "Import a CSV/XLSX file (registry extract, member list, own sheet) with automatic column mapping."
    )
    trust = 65

    def __init__(
        self,
        ctx,
        path: Path,
        source_label: str = "csv_import",
        mapping: dict[str, str] | None = None,
        kind: SourceKind = SourceKind.IMPORT,
    ):
        super().__init__(ctx)
        self.path = Path(path)
        self.source_label = source_label
        self.explicit_mapping = mapping
        self.record_kind = kind

    async def search(self, spec: SourceSearchSpec) -> AsyncIterator[RawRecord]:
        headers, rows = read_rows(self.path)
        mapping = {**auto_map(headers), **(self.explicit_mapping or {})}
        if "name" not in mapping:
            raise ValueError(f"Could not find a business-name column in {self.path.name}; headers={headers}")
        for row in rows:
            rec = row_to_record(row, mapping, self.source_label, spec.naics, self.record_kind)
            if rec:
                yield rec
